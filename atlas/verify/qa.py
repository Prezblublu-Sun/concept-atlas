"""Blind review of LLM drafts and the accept gate of `terms` / `classify` (increment 4 §2.3).

`atlas qa sample --stage terms|classify --n 50 --seed s [--run id]` draws a seeded sample of a draft's items
into `<project>/evalset/<stage>_<run>.xlsx` (sheet `review`: key, item, proposal, evidence sentences cleaned,
judgement — empty — and note; sheet `meta` with task, prompt sha, scheme sha, run, stage, seed) and the same rows as
`<stage>_<run>.jsonl`, each keyed by (task, prompt_sha, scheme_sha, run, item). The reviewer fills `judgement`
(1/0, y/n, yes/no, true/false, ok/bad, 对/错, 是/否, ✓/✗). `atlas qa record <xlsx>` reads the sheet back into
`evalset/judgments/<stage>_<run>.jsonl`; `atlas qa report` aggregates accuracy per stage and stratum (terms: term
type; classify: domain) over every judgment file; `accept_status(project, task, prompt_sha)` is what the accept gate
reads: passed when at least `qa.min_judged` (default 10) items of that task + prompt sha were judged and the
accuracy is ≥ `qa.accept_threshold` (default 0.85).
"""
from __future__ import annotations

import json
import random
from collections.abc import Mapping, Sequence
from pathlib import Path

from atlas.config import Project
from atlas.pipeline.llm_stage import clean_text, drafts_dir, list_drafts, read_draft
from atlas.store.evidence import EvidenceStore

EVALSET_DIR = "evalset"
JUDGMENTS_DIR = "judgments"
STAGE_TASK = {"terms": "terms_extract", "classify": "classify"}
SHEET = "review"
META_SHEET = "meta"
COLUMNS = ("key", "item", "proposal", "evidence", "judgement", "note")
TRUE_WORDS = frozenset({"1", "y", "yes", "true", "t", "ok", "good", "accept", "对", "是", "正确", "✓", "√", "pass"})
FALSE_WORDS = frozenset({"0", "n", "no", "false", "f", "bad", "wrong", "reject", "错", "否", "错误", "✗", "×", "x", "fail"})
DEFAULT_THRESHOLD = 0.85
DEFAULT_MIN_JUDGED = 10


def evalset_dir(project: Project) -> Path:
    return project.root / EVALSET_DIR


def _draft_for(project: Project, stage: str, run_id: str | None) -> Path:
    if stage not in STAGE_TASK:
        raise ValueError(f"qa sample supports stages {sorted(STAGE_TASK)}, not {stage!r}")
    if run_id:
        for name in (f"{run_id}.yaml", f"{run_id}.delta.yaml"):
            p = drafts_dir(project, stage) / name
            if p.is_file():
                return p
        raise FileNotFoundError(f"no {stage} draft for run {run_id}")
    ds = list_drafts(project, stage)
    if not ds:
        raise FileNotFoundError(f"no {stage} draft in {drafts_dir(project, stage)}; run `atlas run {stage} … --pilot N` first")
    return ds[-1]


def parse_judgement(value) -> bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value) if value in (0, 1) else None
    s = str(value).strip().lower()
    if not s:
        return None
    if s in TRUE_WORDS:
        return True
    if s in FALSE_WORDS:
        return False
    return None


def _rows_for(project: Project, stage: str, draft: Mapping) -> list[dict]:
    """One review row per draft item: {key, item, stratum, proposal, sids}."""
    rows: list[dict] = []
    if stage == "terms":
        for t in draft.get("terms") or []:
            rows.append({"key": str(t.get("key") or t.get("term")), "item": str(t.get("term")), "stratum": str(t.get("type") or "?"),
                         "proposal": f"{t.get('zh') or '-'} | {t.get('type')} | {t.get('gloss') or ''} | votes {t.get('votes')}",
                         "sids": list(t.get("example_sids") or [])})
    else:
        for r in draft.get("rows") or []:
            if r.get("status") not in (None, "ok", "skipped"):
                continue
            rows.append({"key": str(r["concept_id"]), "item": str(r["concept_id"]), "stratum": str(r.get("domain") or "?"),
                         "proposal": f"{r.get('domain')} | {r.get('primary')} | {r.get('secondary') or '-'} | {r.get('reason') or ''}",
                         "sids": list(r.get("evidence_sids") or [])})
    return rows


def sample(project: Project, stage: str, n: int = 50, seed: int = 1, run_id: str | None = None) -> dict:
    from openpyxl import Workbook
    p = _draft_for(project, stage, run_id)
    draft = read_draft(p)
    rows = sorted(_rows_for(project, stage, draft), key=lambda r: r["key"])
    rng = random.Random(int(seed))
    chosen = sorted(rng.sample(rows, min(int(n), len(rows))), key=lambda r: r["key"]) if rows else []
    sids = [s for r in chosen for s in r["sids"]]
    texts: dict[str, str] = {}
    if sids and project.evidence_db.is_file():
        store = EvidenceStore.open(project.evidence_db, "ro")
        try:
            texts = store.texts_of(sids)
        finally:
            store.close()
    run = str(draft.get("run_id") or p.stem.split(".")[0])
    meta = {"task": str(draft.get("task") or STAGE_TASK[stage]), "prompt_sha": str(draft.get("prompt_sha") or ""),
            "scheme_sha": str(draft.get("scheme_sha") or "-"), "run": run, "stage": stage, "seed": int(seed),
            "draft": str(p), "n": len(chosen)}
    out_dir = evalset_dir(project)
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"{stage}_{run}.xlsx"
    jsonl = out_dir / f"{stage}_{run}.jsonl"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(list(COLUMNS))
    with open(jsonl, "w", encoding="utf-8") as f:
        for r in chosen:
            evidence = [clean_text(texts[s]) for s in r["sids"] if texts.get(s)]
            ws.append([r["key"], r["item"], r["proposal"], "\n".join(evidence), None, None])
            f.write(json.dumps({**{k: meta[k] for k in ("task", "prompt_sha", "scheme_sha", "run")}, "item": r["item"],
                                "key": r["key"], "stratum": r["stratum"], "proposal": r["proposal"], "evidence": evidence,
                                "sids": r["sids"]}, ensure_ascii=False) + "\n")
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 90
    ms = wb.create_sheet(META_SHEET)
    for k, v in meta.items():
        ms.append([k, v])
    wb.save(xlsx)
    return {"xlsx": str(xlsx), "jsonl": str(jsonl), "n": len(chosen), "total": len(rows), **meta}


def record(project: Project, xlsx: str | Path) -> dict:
    from openpyxl import load_workbook
    p = Path(xlsx)
    if not p.is_file():
        raise FileNotFoundError(f"sheet not found: {p}")
    wb = load_workbook(p, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames or META_SHEET not in wb.sheetnames:
        raise ValueError(f"{p.name}: not a qa sheet (needs sheets {SHEET!r} and {META_SHEET!r})")
    meta = {str(r[0]): r[1] for r in wb[META_SHEET].iter_rows(values_only=True) if r and r[0] is not None}
    stage, run = str(meta.get("stage")), str(meta.get("run"))
    strata: dict[str, str] = {}
    src = evalset_dir(project) / f"{stage}_{run}.jsonl"
    if src.is_file():
        with open(src, encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    d = json.loads(line)
                    strata[str(d["key"])] = str(d.get("stratum") or "?")
    rows = []
    unparsed = []
    it = wb[SHEET].iter_rows(values_only=True)
    header = [str(h) for h in next(it, ())]
    idx = {h: i for i, h in enumerate(header)}
    for r in it:
        if not r or r[idx["key"]] is None:
            continue
        key = str(r[idx["key"]])
        j = parse_judgement(r[idx["judgement"]] if "judgement" in idx else None)
        raw = r[idx["judgement"]] if "judgement" in idx else None
        if j is None and raw not in (None, ""):
            unparsed.append({"key": key, "value": raw})
        rows.append({"task": str(meta.get("task")), "prompt_sha": str(meta.get("prompt_sha")),
                     "scheme_sha": str(meta.get("scheme_sha") or "-"), "run": run, "stage": stage, "item": str(r[idx["item"]]),
                     "key": key, "stratum": strata.get(key, "?"), "judgement": j,
                     "note": str(r[idx["note"]]) if "note" in idx and r[idx["note"]] is not None else ""})
    wb.close()
    out_dir = evalset_dir(project) / JUDGMENTS_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{stage}_{run}.jsonl"
    with open(out, "w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
    return {"jsonl": str(out), "n": len(rows), "judged": sum(1 for x in rows if x["judgement"] is not None),
            "correct": sum(1 for x in rows if x["judgement"] is True), "unparsed": unparsed, "stage": stage, "run": run,
            "task": str(meta.get("task")), "prompt_sha": str(meta.get("prompt_sha"))}


def _judgments(project: Project) -> list[dict]:
    d = evalset_dir(project) / JUDGMENTS_DIR
    rows: list[dict] = []
    if not d.is_dir():
        return rows
    for p in sorted(d.glob("*.jsonl")):
        with open(p, encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    rows.append(json.loads(line))
    return rows


def _acc(rows: Sequence[Mapping]) -> dict:
    judged = [r for r in rows if r.get("judgement") is not None]
    correct = sum(1 for r in judged if r["judgement"] is True)
    return {"n": len(rows), "judged": len(judged), "correct": correct,
            "accuracy": round(correct / len(judged), 4) if judged else None}


def thresholds(project: Project) -> tuple[float, int]:
    thr = project.get("qa.accept_threshold", DEFAULT_THRESHOLD)
    mn = project.get("qa.min_judged", DEFAULT_MIN_JUDGED)
    return float(DEFAULT_THRESHOLD if thr is None else thr), int(DEFAULT_MIN_JUDGED if mn is None else mn)


def report(project: Project) -> dict:
    """{threshold, min_judged, stages: {stage: {task, prompts: {prompt_sha: {runs, n, judged, correct, accuracy,
    passed, strata: {stratum: {...}}}}}}} over every recorded judgment."""
    thr, mn = thresholds(project)
    out: dict = {"threshold": thr, "min_judged": mn, "stages": {}}
    by: dict[tuple[str, str], list[dict]] = {}
    for r in _judgments(project):
        by.setdefault((str(r.get("stage")), str(r.get("prompt_sha"))), []).append(r)
    for (stage, sha), rows in sorted(by.items()):
        st = out["stages"].setdefault(stage, {"task": rows[0].get("task"), "prompts": {}})
        a = _acc(rows)
        strata: dict[str, list[dict]] = {}
        for r in rows:
            strata.setdefault(str(r.get("stratum") or "?"), []).append(r)
        st["prompts"][sha] = {**a, "runs": sorted({str(r.get("run")) for r in rows}),
                              "passed": bool(a["judged"] >= mn and a["accuracy"] is not None and a["accuracy"] >= thr),
                              "strata": {k: _acc(v) for k, v in sorted(strata.items())}}
    return out


def accept_status(project: Project, task: str, prompt_sha: str) -> dict:
    """What the accept gate reads: {passed, accuracy, n, judged, threshold, min_judged, reason}."""
    thr, mn = thresholds(project)
    rows = [r for r in _judgments(project) if str(r.get("task")) == task and str(r.get("prompt_sha")) == prompt_sha]
    a = _acc(rows)
    if not rows:
        reason = "no blind review recorded"
    elif a["judged"] < mn:
        reason = f"{a['judged']} judged < min_judged {mn}"
    elif a["accuracy"] < thr:
        reason = f"accuracy {a['accuracy']} < threshold {thr}"
    else:
        reason = "passed"
    return {"passed": reason == "passed", "accuracy": a["accuracy"], "n": a["n"], "judged": a["judged"], "threshold": thr,
            "min_judged": mn, "reason": reason}


def format_report(r: Mapping) -> str:
    lines = [f"qa report (threshold {r['threshold']}, min judged {r['min_judged']})"]
    if not r["stages"]:
        lines.append("  no judgments recorded (atlas qa sample → fill the sheet → atlas qa record <xlsx>)")
    for stage, st in r["stages"].items():
        for sha, p in st["prompts"].items():
            lines.append(f"  {stage} ({st['task']} @ {sha[:12]}): judged {p['judged']}/{p['n']}, accuracy "
                         f"{p['accuracy'] if p['accuracy'] is not None else '-'}, {'PASSED' if p['passed'] else 'not passed'}; runs {', '.join(p['runs'])}")
            for k, v in p["strata"].items():
                lines.append(f"    {k}: {v['correct']}/{v['judged']} judged of {v['n']}"
                             + (f" ({v['accuracy']})" if v["accuracy"] is not None else ""))
    return "\n".join(lines)
