"""Increment 4 chain on the fixture corpus with a FakeChatClient that answers from the request payloads (schema-valid,
evidence-bound): terms --pilot → qa → curate merge → concepts → merge + confirm → scheme → merge + confirm ×3 →
build new → match → assign → classify (pilot, qa, full) → merge → finish. Asserts draft and curated shapes,
fingerprints, gate messages, dry-run plans, the no-network guard, batch / budget handling, the CLI verbs, and that
`curated/` is written only by `curate` commands (static + runtime)."""
from __future__ import annotations

import json
import os
import re
import shutil
import stat

import pytest
import yaml
from openpyxl import load_workbook

from atlas import cli
from atlas.config import Project, deep_merge, init_project, load_project
from atlas.core.matcher import DEFAULT_PATTERNS
from atlas.llm.base import NO_NETWORK_ENV, BatchStatus, ChatRequest, ChatResponse, FakeChatClient, NetworkForbidden
from atlas.llm.client import make_client
from atlas.pipeline import build as B
from atlas.pipeline import gates
from atlas.pipeline.llm_stage import drafts_dir, format_plan, read_draft, run_id_of
from atlas.pipeline.runner import StageError, run_stage
from atlas.pipeline.stages.ingest import register_documents
from atlas.store.build import BuildStore, read_manifest
from atlas.store.curated import curated_path, curated_shas, read_curated, write_yaml
from atlas.store.llm import LlmStore
from atlas.tools.fixture import make_fixture
from atlas.verify import qa

LEGACY_CONCEPT_KEYS = {"id", "label", "zh", "type", "mode", "forms", "gen_forms", "case_forms", "rx", "probe", "readable",
                       "domains", "place", "shared_with", "origin", "merged_into", "source"}
PHRASES = {"micromotion": "量与指标", "cementless stem": "名词术语", "primary stability": "量与指标", "bone ingrowth": "名词术语",
           "proximal femur": "名词术语", "cancellous bone": "名词术语", "finite element": "方法或模型名",
           "stress shielding": "名词术语", "load transfer": "句式短语"}
SENT_RE = re.compile(r"^\s*\[(S1[0-9a-f]{16})\] (.*)$", re.M)
TERM_RE = re.compile(r"^- term: (.+?) \| zh: (.*?) \| type: (.+?) \|", re.M)


# ---- canned answers derived from the payloads -------------------------------------------------------------------------

def answer_terms(req):
    terms = []
    seen = set()
    for sid, text in SENT_RE.findall(req.user):
        low = text.lower()
        for phrase, typ in PHRASES.items():
            if phrase in low and phrase not in seen and len(terms) < 24:
                seen.add(phrase)
                terms.append({"term": phrase, "zh": "译名", "type": typ, "gloss": "a short gloss", "sids": [sid]})
    if terms:   # a foreign sid is dropped by the stage; an over-long gloss is trimmed
        terms[0] = dict(terms[0], sids=terms[0]["sids"] + ["S1ffffffffffffffff"], gloss=" ".join(["w"] * 30))
        terms.append({"term": "ghost term", "zh": "", "type": "名词术语", "gloss": "", "sids": ["S1ffffffffffffffff"]})
    return json.dumps({"terms": terms}, ensure_ascii=False)


def answer_concepts(req):
    concepts = []
    for block in req.user.split("## G")[1:]:
        terms = TERM_RE.findall(block)
        sids = SENT_RE.findall(block)
        if not terms:
            continue
        label = terms[0][0]
        forms = [t[0] for t in terms] + [label + "s"]
        if "micromotion" in label and label != "micromotion":
            forms.append("micromotion")          # the head noun as a form: collides with C_micromotion in delta mode
        concepts.append({"id": "C_" + re.sub(r"[^a-z0-9]", "", label.lower()), "label": label,
                         "zh": terms[0][1] if terms[0][1] not in ("", "-") else "译", "type": terms[0][2], "forms": forms,
                         "case_forms": [], "merged_terms": [t[0] for t in terms], "sids": [s for s, _ in sids][:2], "note": ""})
    return json.dumps({"concepts": concepts, "deltas": []}, ensure_ascii=False)


def answer_scheme(req):
    ids = re.findall(r"^(C_[a-z0-9]+): ", req.user, re.M)
    docs = re.findall(r"^\[(P\d{3})\]", req.user, re.M)
    mm = [i for i in ids if "micromotion" in i or "primarystability" in i]
    fe = [i for i in ids if "finiteelement" in i]
    subs = [{"id": "ORT-MM", "kind": "within", "domain": "DOM_ortho", "name_zh": "微动", "name_en": "Micromotion", "desc_zh": "微动相关",
             "desc_en": "about micromotion", "anchors": mm + ["C_nope"], "evidence_docs": docs[:2]},
            {"id": "bad id", "kind": "within", "domain": "DOM_fe", "name_zh": "建模", "name_en": "FE modelling", "desc_zh": "", "desc_en": "",
             "anchors": fe}]
    if fe and mm:
        subs.append({"id": "X-FEMM", "kind": "cross", "domain": None, "name_zh": "有限元×微动", "name_en": "FE x micromotion",
                     "desc_zh": "", "desc_en": "", "groups": [fe, mm[:1]]})
    return json.dumps({
        "domains": [{"id": "DOM_ortho", "zh": "骨科", "en": "Orthopaedics", "scope": True},
                    {"id": "DOM_fe", "zh": "有限元", "en": "Finite element", "scope": True}],
        "categories": [{"id": "O1", "domain": "DOM_ortho", "name_zh": "指标", "name_en": "Metrics", "desc_zh": "量", "desc_en": "quantities"},
                       {"id": "O2", "domain": "DOM_ortho", "name_zh": "结构", "name_en": "Structures", "desc_zh": "", "desc_en": ""},
                       {"id": "F1", "domain": "DOM_fe", "name_zh": "方法", "name_en": "Methods", "desc_zh": "", "desc_en": ""},
                       {"id": "Z9", "domain": "DOM_nope", "name_zh": "x", "name_en": "x", "desc_zh": "", "desc_en": ""}],
        "subdomains": subs}, ensure_ascii=False)


def answer_classify(req):
    cid = re.search(r"^Concept: (\S+)", req.user, re.M).group(1)
    sids = [s for s, _ in SENT_RE.findall(req.user)]
    dom = "DOM_fe" if "finiteelement" in cid else "DOM_ortho"
    primary = "F1" if dom == "DOM_fe" else ("O1" if "micromotion" in cid or "stability" in cid else "O2")
    return json.dumps({"domain": dom, "primary": primary, "secondary": None, "reason": "by the example sentences", "sids": sids[:2]})


def fake_client() -> FakeChatClient:
    return FakeChatClient({"terms_extract": answer_terms, "concepts_merge": answer_concepts, "scheme_draft": answer_scheme,
                           "classify": answer_classify})


def fill_sheet(xlsx: str, value) -> None:
    wb = load_workbook(xlsx)
    ws = wb[qa.SHEET]
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=5, value=value)
    wb.save(xlsx)


def _unlock(project) -> None:
    for p in project.builds_dir.glob("*/build.sqlite"):
        os.chmod(p, stat.S_IWRITE | stat.S_IREAD)


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    base = tmp_path_factory.mktemp("curation 测试")
    info = make_fixture(base / "fx", 1)
    project = init_project(base / "proj", name="curation", pdf_dir=str(base / "fx" / "pdf"))
    with open(project.root / "project.yaml", "a", encoding="utf-8") as f:
        yaml.safe_dump({"llm": {"provider": "fake", "model": "claude-opus-5-5", "batch_default": False},
                        "qa": {"min_judged": 3}, "curation": {"terms": {"sample_n": 300, "per_doc_cap": 60}}}, f, allow_unicode=True)
    project = load_project(project.root)
    shutil.copy(info["manifest"], project.manifest_path)
    assert register_documents(project).items_failed == 0
    assert run_stage(project, "ingest").items_failed == 0
    yield {"info": info, "project": project, "state": {}}
    _unlock(project)


def run(env, stage, **opts):
    """run_stage with the runtime check that no stage touches curated/."""
    project = env["project"]
    before = curated_shas(project.curated_dir)
    rep = run_stage(project, stage, **opts)
    assert curated_shas(project.curated_dir) == before, f"{stage} wrote curated/"
    return rep


# ---- tests (ordered; they build on each other through `env['state']`) -----------------------------------------------------

def test_dry_run_plans_without_any_call_or_file(env):
    project = env["project"]
    fake = fake_client()
    rep = run(env, "terms", dry_run=True, pilot=3, client=fake)
    plan = rep.details["plan"]
    assert rep.details["dry_run"] and fake.requests == [] and not drafts_dir(project, "terms").exists()
    assert plan["stage"] == "terms" and plan["pilot"] == 3 and plan["items"] == rep.items_total > 0
    assert plan["sample"]["documents"] == 3 and plan["sample"]["sentences"] <= 180 and plan["sample"]["chunks"] == plan["items"]
    assert set(plan["sample"]["composition"]) >= {"unit_kind|section", "venue", "year_bin", "section"}
    assert 0 < plan["cost_low"] < plan["cost_high"] and plan["cost_batch_high"] == pytest.approx(plan["cost_high"] / 2, rel=1e-3)
    assert plan["tokens"]["input_high"] >= plan["tokens"]["input_low"] > 0 and plan["model"] == "claude-opus-5-5"
    txt = format_plan(plan)
    assert "cost live:" in txt and "cost batch:" in txt and "USD" in txt and "known cache hits: 0" in txt
    assert all(re.fullmatch(r"terms1:[0-9a-f]{12}:[0-9a-f]{16}", fp) for fp in rep.details["fingerprints"].values())
    full = run(env, "terms", dry_run=True, client=fake)
    assert full.details["plan"]["sample"]["documents"] == 9 and full.details["plan"]["items"] >= plan["items"]
    assert fake.requests == [] and not (project.root / "llm" / "llm.sqlite").exists()
    cheap = run(env, "terms", dry_run=True, pilot=3, model="claude-haiku-5-5", client=fake).details["plan"]
    assert cheap["model"] == "claude-haiku-5-5" and cheap["items"] == plan["items"] and cheap["cost_high"] < plan["cost_high"] / 10
    assert run(env, "terms", dry_run=True, pilot=3, model="no-such-model", client=fake).details["plan"]["model"] == "claude-opus-5-5"


def test_no_network_guard_is_honoured(env, monkeypatch):
    """With ATLAS_LLM_NO_NETWORK set, a non-dry run of an LLM stage on a real provider raises NetworkForbidden before
    any provider call: no SDK object is built, no ledger exists, no draft is written. A dry run on the same provider
    still plans — with the configured split — and creates nothing either."""
    import anthropic
    project = env["project"]
    monkeypatch.setenv(NO_NETWORK_ENV, "1")
    monkeypatch.setattr(anthropic, "Anthropic", lambda *a, **k: pytest.fail("the SDK client was built under the guard"))
    real = Project(root=project.root, config=deep_merge(project.config, {"llm": {"provider": "anthropic"}}))
    with pytest.raises(NetworkForbidden, match="ATLAS_LLM_NO_NETWORK"):
        run_stage(real, "terms", pilot=2)
    assert not drafts_dir(project, "terms").exists() and not project.llm_db.exists()
    rep = run_stage(real, "terms", dry_run=True, pilot=2)          # a dry run never calls and never writes
    plan = rep.details["plan"]
    assert rep.details["dry_run"] and not drafts_dir(project, "terms").exists() and not project.llm_db.exists()
    assert plan["models"] == {"anthropic": "claude-opus-5-5", "deepseek": "deepseek-flash"} and plan["currency"] == "USD"
    assert plan["split"] == {"anthropic": pytest.approx(0.9), "deepseek": pytest.approx(0.1)}
    assert set(plan["providers"]) <= {"anthropic", "deepseek"} and sum(v["items"] for v in plan["providers"].values()) == plan["items"] > 0
    assert all(r["provider"] in ("anthropic", "deepseek") and not r["cached"] for r in plan["rows"])
    assert plan["known_cache_hits"] == 0 and 0 < plan["cost_low"] < plan["cost_high"]
    txt = format_plan(plan)
    assert "anthropic (claude-opus-5-5, share 90%)" in txt and "deepseek (deepseek-flash, share 10%)" in txt
    assert "no batch API: this share runs live" in txt and "no provider call is made by a dry run" in txt
    # the second line of defence: a real client refuses on its own before touching the SDK
    with pytest.raises(NetworkForbidden):
        make_client(real).complete(ChatRequest(model="claude-opus-5-5", system=("s",), user="ping", max_tokens=5))
    assert not (project.root / "work" / "runs").exists()


def test_terms_pilot_draft_and_accept_gate(env):
    project, info = env["project"], env["info"]
    fake = fake_client()
    with pytest.raises(StageError, match="no passing blind review"):
        run(env, "terms", client=fake)
    assert fake.requests == []
    rep = run(env, "terms", pilot=3, seed=1, client=fake)
    assert rep.items_failed == 0 and rep.items_done == rep.items_total == len(fake.requests) > 0
    d = read_draft(rep.details["draft"])
    assert d["stage"] == "terms" and d["version"] == "terms1" and d["task"] == "terms_extract" and d["pilot"] == 3
    assert d["prompt_sha"] == rep.details["prompt_sha"] and len(d["prompt_sha"]) == 64 and d["mode"] == "live"
    assert all(re.fullmatch(r"terms1:[0-9a-f]{12}:[0-9a-f]{16}", it["fingerprint"]) and it["status"] == "ok" for it in d["items"])
    assert all(it["fingerprint"] == rep.details["fingerprints"][it["key"]] for it in d["items"])
    terms = {t["term"]: t for t in d["terms"]}
    assert "micromotion" in terms and "ghost term" not in terms and d["stats"]["dropped_no_evidence"] >= 1
    assert d["stats"]["foreign_sids"] >= 1 and d["stats"]["gloss_trimmed"] >= 1
    assert all(len(t["gloss"].split()) <= 20 for t in d["terms"])
    mm = terms["micromotion"]
    assert mm["type"] == "量与指标" and mm["votes"] >= 1 and set(mm["docs"]) <= {x["doc_id"] for x in info["documents"]}
    assert all(re.fullmatch(r"S1[0-9a-f]{16}", s) for s in mm["example_sids"]) and mm["n_docs"] == len(mm["docs"])
    assert (project.root / "work" / "drafts" / "terms" / f"{d['run_id']}.report.md").read_text(encoding="utf-8").startswith("# terms run")
    assert d["cost"] > 0 and d["usage"]["input"] > 0 and d["items"][0]["usage"]["output"] > 0
    assert all(it["provider"] == "fake" and it["model_reported"] == "claude-opus-5-5" for it in d["items"])
    assert d["ledger_run_id"] == f"terms_{d['run_id']}" and d["provider"] == "fake" and d["split"] == {"fake": 1.0}
    assert d["providers"] == {"fake": {"items": d["n_items"], "cost": pytest.approx(d["cost"], abs=1e-5)}}
    assert os.path.isfile(d["settings"]) and d["settings"].endswith("llm_settings.json") and d["ledger_run_id"] in d["settings"]
    assert json.load(open(d["settings"], encoding="utf-8"))["tasks"]["terms_extract"]["prompt_sha"] == d["prompt_sha"]
    # the ledger has one ok row per item for this run, priced like the draft
    store = LlmStore.open(project.llm_db)
    rows = store.ledger_rows(run_id=d["ledger_run_id"])
    store.close()
    assert len(rows) == d["n_items"] and all(r["status"] == "ok" and r["stage"] == "terms" and r["provider"] == "fake" for r in rows)
    assert sum(r["cost"] for r in rows) == pytest.approx(d["cost"], abs=1e-5)
    # the request payloads are pure: sorted sids, no doc ids, every sid of the chunk
    for req in fake.requests:
        sids = [s for s, _ in SENT_RE.findall(req.user)]
        assert sids == sorted(sids) and "P00" not in req.user and req.output_schema["required"] == ["terms"]
        assert req.system[0].startswith("You are a bilingual") and req.key in rep.details["fingerprints"]
    fake2 = fake_client()
    rep2 = run(env, "terms", pilot=3, seed=1, client=fake2)
    d2 = read_draft(rep2.details["draft"])
    assert [it["fingerprint"] for it in d2["items"]] == [it["fingerprint"] for it in d["items"]]     # deterministic
    assert d2["run_id"] != d["run_id"]
    # … and the same items are exact-hash cache hits: answered from llm.sqlite, no request, no cost
    assert fake2.requests == [] and rep2.items_done == rep2.items_total and d2["cost"] == 0
    assert all(it["status"] == "cache_local" and it["cost"] == 0 for it in d2["items"]) and d2["terms"] == d["terms"]
    env["state"]["terms_draft"] = rep.details["draft"]
    # blind review: a failing sheet keeps the gate shut, a passing one opens it
    s = qa.sample(project, "terms", n=4, seed=1)
    assert s["prompt_sha"] == d["prompt_sha"] and os.path.isfile(s["xlsx"])
    fill_sheet(s["xlsx"], "n")
    qa.record(project, s["xlsx"])
    with pytest.raises(StageError, match="accuracy"):
        run(env, "terms", client=fake_client())
    fill_sheet(s["xlsx"], "y")
    qa.record(project, s["xlsx"])
    assert qa.accept_status(project, "terms_extract", d["prompt_sha"])["passed"]
    rep3 = run(env, "terms", client=fake_client())
    assert rep3.items_failed == 0 and read_draft(rep3.details["draft"])["pilot"] is None
    assert read_draft(rep3.details["draft"])["sample"]["documents"] == 9


def test_batch_mode_budget_and_limit(env):
    project = env["project"]

    class SlowBatch(FakeChatClient):
        polls = 0

        def batch_status(self, batch_id):
            self.polls += 1
            st = super().batch_status(batch_id)
            return BatchStatus(batch_id, "in_progress", st.n_requests, 0, 0) if self.polls == 1 else st

    slow = SlowBatch({"terms_extract": answer_terms})
    assert run(env, "terms", dry_run=True, pilot=2, seed=2, client=slow).details["plan"]["known_cache_hits"] == 0
    rep = run(env, "terms", pilot=2, seed=2, mode="batch", client=slow)
    assert rep.details["pending"] and rep.details["batch_id"].startswith("fakebatch") and "draft" not in rep.details
    pending = list(drafts_dir(project, "terms").glob("*.pending.yaml"))
    assert len(pending) == 1 and read_draft(pending[0])["batch_ids"] == [rep.details["batch_id"]]
    store = LlmStore.open(project.llm_db)
    rows = store.ledger_rows(run_id=rep.details["ledger_run_id"])
    assert len(rows) == rep.items_total and all(r["status"] == "pending" and r["batch_id"] == rep.details["batch_id"] for r in rows)
    assert store.get_batch(rep.details["batch_id"])["stage"] == "terms"
    store.close()
    rep2 = run(env, "terms", pilot=2, seed=2, mode="batch", wait=True, client=slow)
    assert rep2.details["batch_id"] == rep.details["batch_id"] and rep2.items_done == rep2.items_total
    assert not list(drafts_dir(project, "terms").glob("*.pending.yaml")) and len(slow.requests) == rep2.items_total
    d = read_draft(rep2.details["draft"])
    assert d["mode"] == "batch" and d["batch_id"] == rep.details["batch_id"] and d["n_ok"] == rep2.items_total
    assert d["ledger_run_id"] == rep.details["ledger_run_id"] and d["run_id"] == run_id_of(pending[0])
    # the collected batch rows are billed at half price and the rerun that read them is a local cache hit
    assert d["cost"] == pytest.approx(sum(it["cost"] for it in d["items"]), abs=1e-5) and d["cost"] > 0
    for it in d["items"]:
        u = it["usage"]
        assert it["status"] == "cache_local" and u["output"] > 0
        assert it["cost"] == pytest.approx((u["input"] * 4 + u["cache_write"] * 5 + u["cache_read"] * 0.2 + u["output"] * 20) / 1e6 * 0.5, abs=1e-6)
    store = LlmStore.open(project.llm_db)
    statuses = sorted(r["status"] for r in store.ledger_rows(run_id=d["ledger_run_id"]))
    store.close()
    assert statuses == ["cache_batch"] * d["n_items"] + ["cache_local"] * d["n_items"]
    live = read_draft(env["state"]["terms_draft"])
    assert d["cost"] / d["n_items"] < sum(it["cost"] for it in live["items"]) / live["n_items"]
    # items already answered live are cache hits in batch mode too: nothing is submitted
    fake3 = fake_client()
    rep3 = run(env, "terms", pilot=3, mode="batch", wait=True, client=fake3)
    d3 = read_draft(rep3.details["draft"])
    assert rep3.items_done == rep3.items_total and d3["mode"] == "batch" and d3["batch_id"] is None and d3["batch_ids"] == []
    assert fake3.requests == [] and fake3.batches == {} and d3["cost"] == 0 and all(it["status"] == "cache_local" for it in d3["items"])
    with pytest.raises(StageError, match="unknown mode"):
        run(env, "terms", pilot=2, mode="nope", client=fake_client())
    # budget: the high estimate of one item is a few cents; a budget below two items leaves items out (uncached
    # items — 4 pilot documents — so that the high estimates are not zero)
    plan = run(env, "terms", dry_run=True, pilot=4, client=fake_client()).details["plan"]
    assert plan["known_cache_hits"] == 0 and plan["per_item_high"] > 0
    rep4 = run(env, "terms", pilot=4, budget=plan["per_item_high"] * 1.5, client=fake_client())
    d4 = read_draft(rep4.details["draft"])
    assert d4["n_items"] == 1 and len(d4["budget_stop"]) == plan["items"] - 1 and any("budget" in w for w in rep4.warnings)
    assert [it["status"] for it in d4["items"]] == ["ok"] and rep4.items_done == 1
    rep5 = run(env, "terms", pilot=3, limit=2, client=fake_client())
    assert rep5.items_total == 2 and read_draft(rep5.details["draft"])["n_items"] == 2
    # a refused item is counted as failed, the rest of the draft stands
    first = sorted(run_stage(project, "terms", dry_run=True, pilot=5, client=fake_client()).details["fingerprints"])[0]
    bad = FakeChatClient({"terms_extract": answer_terms, first: lambda r: ChatResponse(text="", stop_reason="refusal")})
    rep6 = run(env, "terms", pilot=5, client=bad)
    assert rep6.items_failed == 1 and rep6.items_done == rep6.items_total - 1 and any("refused" in w for w in rep6.warnings)
    d6 = read_draft(rep6.details["draft"])
    assert sum(1 for it in d6["items"] if it["status"] == "refused") == 1 and "## Failed items" in open(rep6.details["report"], encoding="utf-8").read()
    store = LlmStore.open(project.llm_db)
    refused = [r for r in store.ledger_rows(run_id=d6["ledger_run_id"]) if r["status"] == "refused"]
    store.close()
    assert len(refused) == 1 and refused[0]["item_key"] == first and refused[0]["cost"] > 0     # a refusal is billed, never retried
    bad2 = FakeChatClient({"terms_extract": answer_terms, first: lambda r: ChatResponse(text="", stop_reason="refusal")})
    rep7 = run(env, "terms", pilot=5, client=bad2)
    assert rep7.items_failed == 1 and [q.key for q in bad2.requests] == [first]                 # a refusal is not cached


def test_merge_terms_then_concepts(env):
    project = env["project"]
    with pytest.raises(StageError, match="needs a merged terms draft"):
        run(env, "concepts", client=fake_client())
    r = gates.merge_draft(project, env["state"]["terms_draft"])
    assert r["files"]["term_tree"]["added"] and read_curated(project.curated_dir, "term_tree")["domains"][0]["id"] == "DOM_curation"
    entries = [t for d in read_curated(project.curated_dir, "term_tree")["domains"] for s in d["sections"] for sub in s["subs"] for t in sub["terms"]]
    assert all(t["source"] == "llm" and t["concept_id"] is None and t["example"] for t in entries)
    fake = fake_client()
    rep = run(env, "concepts", client=fake)
    assert rep.items_failed == 0 and rep.items_done >= 1
    d = read_draft(rep.details["draft"])
    assert d["stage"] == "concepts" and d["mode"] == "base" and d["patterns"] == dict(DEFAULT_PATTERNS)
    assert rep.details["draft"].endswith(f"{d['run_id']}.yaml")
    assert d["inputs"]["candidates"] == len(entries) and d["inputs"]["groups"] >= 1
    assert all(set(c) == LEGACY_CONCEPT_KEYS | {"added_in", "note", "merged_terms", "example_sids"} for c in d["concepts"])
    ids = {c["id"] for c in d["concepts"]}
    assert "C_micromotion" in ids and all(re.fullmatch(r"C_[a-z0-9]+", i) for i in ids)
    mm = next(c for c in d["concepts"] if c["id"] == "C_micromotion")
    assert "micromotion" in mm["forms"] and mm["source"] == "llm" and mm["added_in"] == d["run_id"] and mm["example_sids"]
    assert mm["origin"].startswith("概念") and mm["mode"] == "forms" and mm["domains"] == []
    lt = next((c for c in d["concepts"] if c["label"] == "load transfer"), None)
    assert lt is None or lt["origin"] == "句式短语"
    # every request is evidence-bound and pure
    for req in fake.requests:
        assert "## G1" in req.user and "No existing concepts" in req.user and req.output_schema["required"] == ["concepts"]
    # no duplicate forms across the draft concepts
    keys = [f.lower() for c in d["concepts"] for f in c["forms"]]
    assert len(keys) == len(set(keys))
    r = gates.merge_draft(project, rep.details["draft"])
    assert set(r["files"]) == {"concepts", "patterns"} and len(r["files"]["concepts"]["added"]) == len(d["concepts"])
    cs = read_curated(project.curated_dir, "concepts")["concepts"]
    assert all(set(c) == LEGACY_CONCEPT_KEYS | {"added_in"} for c in cs)
    assert read_curated(project.curated_dir, "patterns") == dict(DEFAULT_PATTERNS)
    gates.confirm(project, "concepts")
    gates.confirm(project, "patterns")
    # delta mode: a merged variant is covered by rule, a genuinely new term becomes an `add`
    tree = read_curated(project.curated_dir, "term_tree")
    sub = tree["domains"][0]["sections"][0]["subs"][0]
    sub["terms"] += [{"term": "micromotions", "zh": "", "type": "量与指标", "concept_id": None, "note": "", "example": "", "example_doc": "", "source": "owner"},
                     {"term": "cortical bone", "zh": "皮质骨", "type": "名词术语", "concept_id": None, "note": "", "example": "The cortical shell was modelled.", "example_doc": "P001", "source": "owner"}]
    write_yaml(curated_path(project.curated_dir, "term_tree"), tree)
    rep2 = run(env, "concepts", client=fake_client())
    d2 = read_draft(rep2.details["draft"])
    assert rep2.details["draft"].endswith(".delta.yaml") and d2["mode"] == "delta"
    assert any(c["term"] == "micromotions" and c["concept_id"] == "C_micromotion" for c in d2["already_covered"])
    assert [c["label"] for c in d2["concepts"]] == ["cortical bone"] and d2["inputs"]["existing_concepts"] == len(cs)
    r2 = gates.merge_draft(project, rep2.details["draft"])
    assert r2["files"]["concepts"]["added"] == ["C_corticalbone"]
    gates.confirm(project, "concepts")
    env["state"]["concepts"] = [c["id"] for c in read_curated(project.curated_dir, "concepts")["concepts"]]
    # in delta mode the payload lists the overlapping existing concepts
    fake3 = fake_client()
    tree = read_curated(project.curated_dir, "term_tree")
    tree["domains"][0]["sections"][0]["subs"][0]["terms"].append(
        {"term": "stem micromotion", "zh": "", "type": "量与指标", "concept_id": None, "note": "", "example": "", "example_doc": "", "source": "owner"})
    write_yaml(curated_path(project.curated_dir, "term_tree"), tree)
    rep3 = run(env, "concepts", client=fake3)
    assert "Existing confirmed concepts" in fake3.requests[0].user and "C_micromotion:" in fake3.requests[0].user
    d3 = read_draft(rep3.details["draft"])
    assert d3["concepts"] == [] and [x["action"] for x in d3["deltas"]] == ["merge_into"] and d3["deltas"][0]["concept_id"] == "C_micromotion"
    r3 = gates.merge_draft(project, rep3.details["draft"])
    assert r3["files"]["concepts"]["deltas_applied"] == [{"action": "new_form", "concept_id": "C_micromotion", "form": "stem micromotion"}]
    gates.confirm(project, "concepts")


def test_scheme_draft_merge_and_confirm(env):
    project, info = env["project"], env["info"]
    fake = fake_client()
    rep = run(env, "scheme", client=fake)
    assert rep.items_total == rep.items_done == 1 and len(fake.requests) == 1
    req = fake.requests[0]
    assert req.key == "scheme" and "[P001]" in req.user and "C_micromotion:" in req.user and req.output_schema["required"] == ["domains", "categories", "subdomains"]
    d = read_draft(rep.details["draft"])
    assert d["inputs"]["freq_source"].startswith("terms draft") and d["inputs"]["documents"] == 9
    assert [x["id"] for x in d["domains"]] == ["DOM_ortho", "DOM_fe"] and all(x["order"] == i for i, x in enumerate(d["domains"]))
    assert {c["id"] for c in d["categories"]} == {"O1", "O2", "F1"} and d["stats"]["categories_unknown_domain"] == 1
    subs = {s["id"]: s for s in d["subdomains"]}
    assert "ORT-MM" in subs and "C_nope" not in subs["ORT-MM"]["anchors"] and d["stats"]["unresolved_anchors"] == ["C_nope"]
    assert subs["ORT-MM"]["evidence_docs"] == ["P001", "P002"] and all(a in env["state"]["concepts"] for a in subs["ORT-MM"]["anchors"])
    assert any(s["kind"] == "within" and s["domain"] == "DOM_fe" and s["id"] == "FE-FM" for s in d["subdomains"])   # 'bad id' regenerated
    if "X-FEMM" in subs:
        assert len(subs["X-FEMM"]["groups"]) == 2 and subs["X-FEMM"]["domain"] is None
    assert all(set(x) == {"id", "zh", "en", "order", "scope", "source", "added_in"} for x in d["domains"])
    assert all(set(x) == {"id", "domain", "name_zh", "name_en", "desc_zh", "desc_en", "order", "source", "added_in"} for x in d["categories"])
    r = gates.merge_draft(project, rep.details["draft"])
    assert set(r["files"]) == {"domains", "categories", "subdomains"}
    for name in ("domains", "categories", "subdomains"):
        gates.confirm(project, name)
    s0 = read_curated(project.curated_dir, "subdomains")["subdomains"][0]
    assert "evidence_docs" not in s0 and set(s0) >= {"id", "kind", "domain", "name_zh", "name_en", "desc_zh", "desc_en", "order", "notes", "anchors"}
    write_yaml(curated_path(project.curated_dir, "doc_domains"), {"docs": {x["doc_id"]: ["DOM_ortho", "DOM_fe"] for x in info["documents"]}})
    gates.confirm(project, "doc_domains")
    notices = gates.draft_notices(project)
    assert not any("`domains`" in n for n in notices)


def test_build_classify_qa_merge_finish(env):
    project = env["project"]
    m = B.new_build(project, note="curation")
    bid = m["build_id"]
    assert run(env, "match", build=bid, workers=1).items_failed == 0
    with pytest.raises(StageError, match="requires `assign`"):
        run(env, "classify", build=bid, pilot=2, client=fake_client())
    assert run(env, "assign", build=bid).items_failed == 0
    fake = fake_client()
    with pytest.raises(StageError, match="no passing blind review"):
        run(env, "classify", build=bid, client=fake)
    rep = run(env, "classify", build=bid, pilot=3, client=fake)
    assert rep.items_total == rep.items_done == 3 == len(fake.requests)
    req = fake.requests[0]
    assert len(req.system) == 2 and req.system[1].startswith("Classification scheme") and "O1:" in req.system[1]
    assert req.output_schema["properties"]["domain"]["enum"] == ["DOM_ortho", "DOM_fe"] and set(req.output_schema["properties"]["primary"]["enum"]) == {"F1", "O1", "O2"}
    assert re.search(r"^Concept: C_", req.user, re.M) and SENT_RE.search(req.user)
    d = read_draft(rep.details["draft"])
    assert d["stage"] == "classify" and d["build_id"] == bid and len(d["scheme_sha"]) == 64 and len(d["rows"]) == 3
    row = d["rows"][0]
    assert set(row) == {"concept_id", "domain", "primary", "secondary", "reason", "evidence_sids", "source", "added_in", "status"}
    assert row["status"] == "ok" and row["primary"] in ("O1", "O2", "F1") and row["evidence_sids"]
    st = BuildStore.open(project.builds_dir / bid, "ro")
    try:
        items = st.stage_items("classify")
        assert sum(1 for it in items if it["item_key"] != "*") == 3 and st.stage_status("classify") == "done"
        assert all(re.fullmatch(r"classify1:[0-9a-f]{12}:[0-9a-f]{16}", it["fingerprint"]) for it in items if it["item_key"] != "*")
    finally:
        st.close()
    s = qa.sample(project, "classify", n=2, seed=1)
    assert s["scheme_sha"] == d["scheme_sha"] and s["n"] == 2
    fill_sheet(s["xlsx"], "y")
    rec = qa.record(project, s["xlsx"])
    assert rec["judged"] == 2
    with pytest.raises(StageError, match="min_judged"):
        run(env, "classify", build=bid, client=fake_client())
    s2 = qa.sample(project, "classify", n=3, seed=2)        # the same run: its judgment file is replaced
    fill_sheet(s2["xlsx"], "y")
    qa.record(project, s2["xlsx"])
    rep_q = qa.report(project)
    assert rep_q["stages"]["classify"]["prompts"][d["prompt_sha"]]["passed"] and "classify" in qa.format_report(rep_q)
    fake2 = fake_client()
    rep2 = run(env, "classify", build=bid, client=fake2)
    d2 = read_draft(rep2.details["draft"])
    assert rep2.items_skipped == 3 and rep2.items_done == rep2.items_total - 3 == len(fake2.requests)
    # every tiered concept with example sentences gets a row (phrases stay untiered; a concept without occurrences
    # has nothing to cite and is skipped); the pilot's three rows are reused, not re-asked
    assert d2["inputs"]["concepts"] <= len(env["state"]["concepts"]) and "C_loadtransfer" not in {r["concept_id"] for r in d2["rows"]}
    assert len(d2["rows"]) == d2["inputs"]["concepts"] - len(d2["stats"]["no_examples"]) >= 5
    assert sum(1 for r in d2["rows"] if r["status"] == "skipped") == 3
    assert d2["stats"]["by_domain"].get("DOM_ortho", 0) >= 1
    rep3 = run(env, "classify", build=bid, client=fake_client())
    assert rep3.items_skipped == rep3.items_total and rep3.items_done == 0        # every fingerprint unchanged
    r = gates.merge_draft(project, rep2.details["draft"])
    rows = read_curated(project.curated_dir, "concept_categories")["rows"]
    assert len(r["files"]["concept_categories"]["added"]) == len(d2["rows"]) and all(set(x) == {"concept_id", "domain", "primary", "secondary", "reason", "source", "added_in"} for x in rows)
    gates.confirm(project, "concept_categories")
    fin = run(env, "finish", build=bid)
    assert fin.items_failed == 0 and read_manifest(project.builds_dir / bid)["status"] == "finished"
    assert read_manifest(project.builds_dir / bid)["stages"]["classify"]["status"] == "done"
    env["state"]["build"] = bid


def test_cli_verbs(env, capsys):
    project = env["project"]
    root = str(project.root)
    assert cli.main(["run", "terms", root, "--dry-run", "--pilot", "2"]) == 0
    out = capsys.readouterr().out
    assert "[terms] dry run" in out and "cost live:" in out and "cost batch:" in out and "pilot 2 documents" in out
    # provider fake without canned answers: the default client answers "{}", the item fails validation (after the
    # one "return only the JSON object" retry), exit 2; a different model is a different request hash, so the
    # answer cached by the earlier batch run of these items is not used
    assert cli.main(["run", "terms", root, "--pilot", "2", "--limit", "1", "--model", "claude-haiku-5-5"]) == 2
    out = capsys.readouterr().out
    assert "[terms] done 0/1" in out and "invalid" in out and "draft:" in out and "ledger_run_id: terms_r" in out
    lrid = re.search(r"ledger_run_id: (\S+)", out).group(1)
    assert cli.main(["cost", "report", root, "--run", lrid]) == 0
    out = capsys.readouterr().out
    assert out.startswith(f"run {lrid}: stage terms, mode live") and "providers: fake claude-haiku-5-5 effort medium" in out
    assert f"ledger (run {lrid}): 2 rows" in out and "fake\tterms\tterms_extract\t2\t0\t0\t0\t2\t0\t" in out    # two malformed attempts
    assert cli.main(["cost", "report", root]) == 0
    out = capsys.readouterr().out
    assert out.startswith("ledger: ") and "cache: local hit rate" in out and "unknown spend: none" in out
    assert cli.main(["curate", "status", root]) == 0
    out = capsys.readouterr().out
    assert "concepts\tconfirmed" in out and "notice:" in out                        # newer terms drafts than term_tree's confirmation
    assert cli.main(["curate", "merge", root, str(project.root / "nope.yaml")]) == 1
    assert "draft not found" in capsys.readouterr().err
    assert cli.main(["curate", "merge", root, env["state"]["terms_draft"]]) == 0
    assert "merged terms draft" in capsys.readouterr().out
    assert cli.main(["qa", "report", root]) == 0
    out = capsys.readouterr().out
    assert "terms (terms_extract" in out and "PASSED" in out
    assert cli.main(["qa", "sample", root, "--stage", "classify", "--n", "2"]) == 0
    xlsx = re.search(r"qa sheet: (.+?\.xlsx) \(", capsys.readouterr().out).group(1)     # the path may contain spaces
    fill_sheet(xlsx, "y")
    assert cli.main(["qa", "record", root, xlsx]) == 0
    assert "recorded 2 judgements" in capsys.readouterr().out
    assert cli.main(["run", "classify", root, "--dry-run", "--build", env["state"]["build"]]) == 1   # finished build: not open
    assert "not open" in capsys.readouterr().err


def test_static_llm_stages_never_write_curated():
    """The new modules carry no curated writer (the runtime check runs inside `run` above)."""
    from pathlib import Path
    atlas = Path(cli.__file__).resolve().parent
    for py in [atlas / "pipeline" / "llm_stage.py", atlas / "verify" / "qa.py",
               *(atlas / "llm").glob("*.py"), *(atlas / "pipeline" / "stages").glob("*.py")]:
        src = py.read_text(encoding="utf-8")
        assert "write_yaml" not in src and "write_confirmations" not in src and "confirm(" not in src, py
